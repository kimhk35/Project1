"""make_eval_hwpx.py와 같은 점수(같은 시드)로 xlsx와 docx를 만든다 HWPX가 열리지 않을 때 쓰는 대안"""
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from make_eval_hwpx import (ITEMS, OPINIONS, PUBLISHERS, RECOMMEND, SUMMARY_ORDER, generate)


def score_rows(sheet):
    rows = [["평가영역", "평가기준", "항목별 점수"] + PUBLISHERS]
    for i, (area, crit, mx) in enumerate(ITEMS):
        rows.append([area, crit, mx] + [sheet[p][i] for p in range(len(PUBLISHERS))])
    rows.append(["합계", "", 100] + [sum(sheet[p]) for p in range(len(PUBLISHERS))])
    return rows


def summary_rows(sheets, totals):
    rank = {p: i + 1 for i, p in enumerate(sorted(range(len(totals)), key=lambda p: -totals[p]))}
    rows = [["출판사명"] + [f"위원 {k + 1}" for k in range(len(sheets))] + ["총점", "평균", "비고"]]
    for p, name in SUMMARY_ORDER:
        rows.append([name] + [sum(s[p]) for s in sheets]
                    + [totals[p], round(totals[p] / len(sheets), 1),
                       f"{rank[p]}순위 추천" if rank[p] <= 3 else ""])
    return rows


def recommend_rows():
    return [["순위", "출판사명", "추 천 의 견"]] + [[i + 1, n, t] for i, (n, t) in enumerate(RECOMMEND)]


# ------------------------------------------------------------------ xlsx

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GRAY = PatternFill("solid", fgColor="E7E6E6")


def put_table(ws, top, rows, bold_first=True):
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = ws.cell(row=top + r, column=c + 1, value=v)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if r == 0 and bold_first:
                cell.font = Font(bold=True)
                cell.fill = GRAY
    return top + len(rows)


def title(ws, text, width):
    ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=16)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")


def write_xlsx(path, sheets, totals):
    wb = Workbook()
    wb.remove(wb.active)
    ncol = 3 + len(PUBLISHERS)
    for k, sheet in enumerate(sheets):
        ws = wb.create_sheet(f"평가표{k + 1}")
        title(ws, "검정(인정)도서 선정 평가표", ncol)
        ws.cell(row=2, column=ncol - 3, value=f"과 목 : 수학과    위 원 : 평가표 {k + 1}    (인)")
        end = put_table(ws, 3, score_rows(sheet))
        ws.cell(row=end, column=1, value="<종합의견 및 추천의견>").font = Font(bold=True)
        ws.merge_cells(start_row=end, start_column=1, end_row=end, end_column=ncol)
        ws.cell(row=end, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=end + 1, column=1, value=OPINIONS[k]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=end + 1, start_column=1, end_row=end + 1, end_column=ncol)
        ws.row_dimensions[end + 1].height = 90
        ws.row_dimensions[3].height = 32
        for c, w in enumerate([14, 26, 10] + [12] * len(PUBLISHERS)):
            ws.column_dimensions[chr(65 + c)].width = w

    ws = wb.create_sheet("총괄표")
    title(ws, "검·인정도서 선정기준 평가 총괄표", 7)
    ws.cell(row=2, column=1, value="과 목 : 수학")
    put_table(ws, 3, summary_rows(sheets, totals))
    for c, w in enumerate([22, 10, 10, 10, 10, 10, 14]):
        ws.column_dimensions[chr(65 + c)].width = w

    ws = wb.create_sheet("추천의견서")
    title(ws, "추천 검·인정도서 및 추천 의견서", 3)
    ws.cell(row=2, column=1, value="과 목 : 수학")
    put_table(ws, 3, recommend_rows())
    for r in range(4, 7):
        ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 110
    for c, w in enumerate([8, 14, 80]):
        ws.column_dimensions[chr(65 + c)].width = w
    wb.save(path)


# ------------------------------------------------------------------ docx

def shade(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "E7E6E6")
    tc_pr.append(shd)


def doc_table(doc, rows, size=9, left_cols=()):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = t.cell(r, c)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (r > 0 and c in left_cols) else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(v))
            run.font.size = Pt(size)
            if r == 0:
                run.bold = True
                shade(cell)
    return t


def heading(doc, text, first=False):
    if not first:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)


def write_docx(path, sheets, totals):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Mm(15))
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    for k, sheet in enumerate(sheets):
        heading(doc, "검정(인정)도서 선정 평가표", first=k == 0)
        p = doc.add_paragraph(f"과 목 : 수학과      위 원 : 평가표 {k + 1}      (인)")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc_table(doc, score_rows(sheet), size=8)
        doc_table(doc, [["<종합의견 및 추천의견>"], [OPINIONS[k]]], size=9, left_cols=(0,))

    heading(doc, "검·인정도서 선정기준 평가 총괄표")
    doc.add_paragraph("과 목 : 수학")
    doc_table(doc, summary_rows(sheets, totals), size=10)

    heading(doc, "추천 검·인정도서 및 추천 의견서")
    doc.add_paragraph("과 목 : 수학")
    t = doc_table(doc, recommend_rows(), size=10, left_cols=(2,))
    for row in t.rows:
        row.cells[0].width, row.cells[1].width, row.cells[2].width = Mm(15), Mm(35), Mm(210)
    doc.save(path)


if __name__ == "__main__":
    base = sys.argv[1] if len(sys.argv) > 1 else "검정도서_선정_평가표_3종"
    seed = int(sys.argv[2]) if len(sys.argv) > 2 else 2026
    sheets, totals = generate(seed)
    write_xlsx(base + ".xlsx", sheets, totals)
    write_docx(base + ".docx", sheets, totals)
    print("saved", base + ".xlsx", base + ".docx")
