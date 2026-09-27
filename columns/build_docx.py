"""AI Insight Column 원고 빌더 (텍스트 원고 -> DOCX, 선택적으로 PDF)

사용법
  python columns/build_docx.py 원고.txt 결과.docx [--pdf]
  --pdf 는 LibreOffice(soffice)로 같은 폴더에 PDF도 만든다

원고 표기 (한 줄 = 한 문단, 빈 줄은 무시)
  #  제목                 15.5pt 굵게 남색
  @  머리 줄              10pt 회색 (시리즈 표기 — 부제)
  ※  안내 문구            9.5pt 회색
  ## 소제목               12pt 굵게 남색
  >  핵심 문장            10.5pt 굵게 남색 좌우 들여쓰기
  |  표 행                칸은 | 로 나눈다 첫 행은 머리 행
  ?  생각해볼 문제 항목    번호를 굵게 (예: ? 1. 질문)
  ### 참고 문헌 제목       11pt 굵게 남색
  -  참고 문헌 항목        9pt 내어쓰기
  그 밖의 줄은 본문 11pt
  **굵게** 는 본문 안에서 굵게

편집 원칙은 columns/EDITING.md 를 따른다 특히 제목과 소제목을 뺀 모든 문단은 양쪽 정렬
"""
import os, re, subprocess, sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

FONT = 'Malgun Gothic'
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
BLACK = RGBColor(0, 0, 0)
JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY


def set_font(run, size, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(k), FONT)


def para(doc, text, size=11, bold=False, color=BLACK, align=JUSTIFY, before=0, after=4,
         spacing=1.5, left=0, right=0, hanging=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), spacing
    if left or hanging:
        pf.left_indent = Mm(left + hanging)
    if hanging:
        pf.first_line_indent = Mm(-hanging)
    if right:
        pf.right_indent = Mm(right)
    # **굵게** 구간을 나눠 run으로
    for i, chunk in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if chunk:
            set_font(p.add_run(chunk), size, bold or i % 2 == 1, color)
    return p


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 or c == 0 else JUSTIFY
            p.paragraph_format.line_spacing = 1.2
            set_font(p.add_run(text), 9.5, bold=(r == 0), color=NAVY if r == 0 else BLACK)
            if r == 0:
                shade(cell, 'DCE3EE')
    # 첫 칸은 좁게
    widths = [Mm(16)] + [Mm((160 - 16) / (len(rows[0]) - 1))] * (len(rows[0]) - 1)
    for col, w in zip(t.columns, widths):
        col.width = w
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    para(doc, '', after=2)


def build(src, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(25)
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    rows = []
    lines = [ln.rstrip() for ln in open(src, encoding='utf-8')]
    for ln in lines + ['']:
        if ln.startswith('|'):
            rows.append([c.strip() for c in ln.strip('|').split('|')])
            continue
        if rows:
            table(doc, rows)
            rows = []
        if not ln.strip():
            continue
        if ln.startswith('### '):
            para(doc, ln[4:], 11, True, NAVY, WD_ALIGN_PARAGRAPH.LEFT, before=22, after=6, spacing=1.2)
        elif ln.startswith('## '):
            para(doc, ln[3:], 12, True, NAVY, WD_ALIGN_PARAGRAPH.LEFT, before=14, after=6, spacing=1.3)
        elif ln.startswith('# '):
            para(doc, ln[2:], 15.5, True, NAVY, WD_ALIGN_PARAGRAPH.LEFT, after=4, spacing=1.2)
        elif ln.startswith('@ '):
            para(doc, ln[2:], 10, color=GRAY, after=10, spacing=1.3)
        elif ln.startswith('※'):
            para(doc, ln, 9.5, color=GRAY, after=8, spacing=1.4)
        elif ln.startswith('> '):
            para(doc, ln[2:], 10.5, True, NAVY, before=4, after=8, left=8, right=8)
        elif ln.startswith('? '):
            m = re.match(r'(\d+\.)\s*(.*)', ln[2:])
            para(doc, f'**{m.group(1)}** {m.group(2)}' if m else ln[2:], after=8)
        elif ln.startswith('- '):
            para(doc, ln[2:], 9, after=2, spacing=1.25, hanging=7)
        else:
            para(doc, ln)
    doc.save(out)


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if len(args) != 2:
        sys.exit(__doc__)
    src, out = args
    build(src, out)
    print('DOCX', out)
    if '--pdf' in sys.argv:
        outdir = os.path.dirname(os.path.abspath(out))
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', outdir, out],
                       check=True, capture_output=True)
        print('PDF', os.path.splitext(out)[0] + '.pdf')
