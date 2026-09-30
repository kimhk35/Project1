"""export/out/ir.json으로 편집 가능한 DOCX를 만든다
사용법  python3 export/make_docx.py  →  dist/Heuresis_Vol01_창간특집호.docx
"""
import json, os, re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
IR = json.load(open(os.path.join(HERE, 'out', 'ir.json'), encoding='utf-8'))
IMG = os.path.join(HERE, 'out', 'img')
OUT = os.path.join(ROOT, 'dist', 'Heuresis_Vol01_창간특집호.docx')

SERIF, SANS, MONO, DISPLAY = 'Noto Serif KR', 'Noto Sans KR', 'IBM Plex Mono', 'Playfair Display'
INK, INK2, MUTED = RGBColor(0x16, 0x16, 0x1A), RGBColor(0x34, 0x32, 0x2F), RGBColor(0x6F, 0x6A, 0x63)
ACCENT = {'': RGBColor(0xC8, 0x37, 0x2D), 'clinic': RGBColor(0x15, 0x60, 0x6A), 'teacher': RGBColor(0xA9, 0x7A, 0x22)}
FILL = {'sand': 'F1EADF', 'navy': '1B2A41', 'teal': 'E1EEEE', 'ochre': 'F4EAD4', 'line': 'FFFFFF', 'card': 'FFFFFF'}
CONTENT_MM = 170


def set_font(target, name, size=None, bold=None, color=None):
    f = target.font
    f.name = name
    rpr = target.element.get_or_add_rPr() if hasattr(target, 'element') else target._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if fonts.get(qn(a)) is not None: del fonts.attrib[qn(a)]
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        fonts.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color


def style(doc, name, font, size, bold=False, color=INK2, before=0, after=4, line=1.55, align=None, base='Normal', keep=False, spacing=None):
    st = doc.styles[name] if name in [s.name for s in doc.styles] else doc.styles.add_style(name, 1)
    if name != 'Normal':
        st.base_style = doc.styles[base]
    set_font(st, font, size, bold, color)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if align is not None: pf.alignment = align
    pf.keep_with_next = keep
    if spacing is not None:
        rpr = st.element.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(spacing)); rpr.append(sp)
    return st


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hex_fill); tcpr.append(sh)


def borders(cell, **sides):
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        sz, col = sides.get(side, (0, 'FFFFFF'))
        el.set(qn('w:val'), 'single' if sz else 'nil'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), col)
        b.append(el)
    tcpr.append(b)


def cell_margins(table, mm=3.2):
    tblpr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}'); el.set(qn('w:w'), str(int(mm * 56.7))); el.set(qn('w:type'), 'dxa'); m.append(el)
    tblpr.append(m)


def full_width(table):
    tblpr = table._tbl.tblPr
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), '5000'); w.set(qn('w:type'), 'pct'); tblpr.append(w)


class Builder:
    def __init__(self):
        self.doc = Document()
        d = self.doc
        style(d, 'Normal', SERIF, 10, color=INK2, after=5, line=1.62, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        for lvl, (font, size, before, after) in {1: (SERIF, 22, 14, 8), 2: (SERIF, 16, 12, 6), 3: (SANS, 11.5, 10, 3), 4: (SANS, 11, 6, 3)}.items():
            st = style(d, f'Heading {lvl}', font, size, True, INK, before, after, 1.3, WD_ALIGN_PARAGRAPH.LEFT, keep=True)
            st.paragraph_format.outline_level = lvl - 1 if hasattr(st.paragraph_format, 'outline_level') else None
        style(d, 'Kicker', MONO, 7.5, color=ACCENT[''], before=6, after=2, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True, spacing=24)
        style(d, 'Dek', SANS, 12, color=INK2, after=8, line=1.55, align=WD_ALIGN_PARAGRAPH.LEFT)
        style(d, 'Small', SANS, 8, color=MUTED, after=4, line=1.4, align=WD_ALIGN_PARAGRAPH.LEFT)
        style(d, 'Quote', SERIF, 13, True, INK, 8, 2, 1.45, WD_ALIGN_PARAGRAPH.LEFT)
        style(d, 'Cite', MONO, 7.5, color=MUTED, after=10, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
        style(d, 'Box Text', SANS, 9, color=INK2, after=3, line=1.55, align=WD_ALIGN_PARAGRAPH.LEFT)
        style(d, 'Table Text', SANS, 8.5, color=INK2, after=0, line=1.35, align=WD_ALIGN_PARAGRAPH.LEFT)
        style(d, 'Caption Img', SANS, 7.5, color=MUTED, after=6, line=1.3, align=WD_ALIGN_PARAGRAPH.LEFT)
        self.theme = ''
        # 표지 구역  여백 없이 한 쪽 전체를 그림으로
        s = d.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin', 'header_distance', 'footer_distance'):
            setattr(s, side, Mm(0))

    # 글자 단위
    def add_runs(self, p, runs, box_dark=False, force_bold=None):
        for r in runs:
            parts = r['t'].split('\n')
            for k, part in enumerate(parts):
                if k: p.add_run().add_break(WD_BREAK.LINE)
                if not part: continue
                run = p.add_run(part)
                if r.get('b') or force_bold: run.bold = True
                if r.get('i'): run.italic = True
                if r.get('m'): set_font(run, MONO)
                if box_dark: run.font.color.rgb = RGBColor(0xEE, 0xF0, 0xF4)
        return p

    def para(self, container, style_name, runs=None, text=None, **kw):
        p = container.add_paragraph(style=style_name)
        if runs is not None: self.add_runs(p, runs, **kw)
        elif text: p.add_run(text)
        return p

    def image(self, container, b, width_mm=None):
        path = os.path.join(IMG, b['id'] + '.png')
        if not os.path.exists(path): return
        w = width_mm or min(b['w'] * 0.2646, CONTENT_MM)
        if b['kind'] == 'band': w = CONTENT_MM
        p = container.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(3), Pt(5)
        p.add_run().add_picture(path, width=Mm(w))

    def table(self, container, b):
        rows = b['rows']
        ncol = max(len(r) for r in rows)
        t = container.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        full_width(t); cell_margins(t, 1.4)
        for i, row in enumerate(rows):
            for j in range(ncol):
                cell = t.cell(i, j)
                c = row[j] if j < len(row) else {'r': [], 'th': False}
                p = cell.paragraphs[0]; p.style = self.doc.styles['Table Text']
                self.add_runs(p, c['r'], force_bold=c['th'] or None)
                top = (12, '16161A') if i == 0 else (0, 'FFFFFF')
                bottom = (8, '16161A') if c['th'] else (4, 'D8D0C2')
                borders(cell, top=top, bottom=bottom)
        container.add_paragraph(style='Small')

    def box(self, container, b):
        t = container.add_table(rows=1, cols=1)
        full_width(t); cell_margins(t, 3.2)
        cell = t.cell(0, 0)
        shade(cell, FILL[b['v']])
        acc = ACCENT[self.theme].__str__() if hasattr(ACCENT[self.theme], '__str__') else 'C8372D'
        if b['v'] == 'line': borders(cell, top=(6, '16161A'), bottom=(6, '16161A'), left=(6, '16161A'), right=(6, '16161A'))
        elif b['v'] == 'card': borders(cell, top=(16, '16161A'), bottom=(4, 'D8D0C2'))
        else: borders(cell)
        dark = b['v'] == 'navy'
        first = cell.paragraphs[0]
        used = False
        for c in b['c']:
            if not used:
                p = first; used = True
            else:
                p = None
            self.block(cell, c, in_box=True, dark=dark, reuse=p)
        container.add_paragraph(style='Small').paragraph_format.space_after = Pt(2)

    def block(self, container, b, in_box=False, dark=False, reuse=None):
        t = b['t']
        def P(style_name):
            if reuse is not None:
                reuse.style = self.doc.styles[style_name]; return reuse
            return container.add_paragraph(style=style_name)
        if t == 'img':
            if reuse is not None: reuse.add_run()
            return self.image(container, b, None)
        if t == 'table':
            return self.table(container, b)
        if t == 'box':
            if in_box:  # 상자 속 상자는 평평하게 편다
                for c in b['c']: self.block(container, c, True, dark)
                return
            return self.box(container, b)
        if t == 'h':
            lvl = b['l']
            p = P('Box Text' if in_box else f'Heading {lvl}')
            self.add_runs(p, b['r'], box_dark=dark, force_bold=True)
            if in_box: p.runs and [setattr(r.font, 'size', Pt(11 if lvl <= 3 else 10)) for r in p.runs]
            return
        if t == 'label':
            p = P('Kicker'); self.add_runs(p, b['r'])
            for r in p.runs: r.font.color.rgb = RGBColor(0xF3, 0xA8, 0x9F) if dark else ACCENT[self.theme]
            return
        if t == 'dek':
            p = P('Dek'); self.add_runs(p, b['r'], box_dark=dark); return
        if t == 'small':
            p = P('Small'); self.add_runs(p, b['r']); return
        if t == 'quote':
            p = P('Quote'); self.add_runs(p, b['r'])
            pbdr = OxmlElement('w:pBdr'); top = OxmlElement('w:top')
            top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '14'); top.set(qn('w:space'), '6'); top.set(qn('w:color'), str(ACCENT[self.theme]))
            pbdr.append(top); p._p.get_or_add_pPr().append(pbdr)
            if b.get('cite'): self.para(container, 'Cite', text=b['cite'])
            return
        if t == 'li':
            p = P('Box Text' if in_box else 'Normal')
            p.paragraph_format.left_indent = Mm(5); p.paragraph_format.first_line_indent = Mm(-4)
            mark = '☐  ' if b.get('check') else '•  '
            p.add_run(mark)
            self.add_runs(p, b['r'], box_dark=dark); return
        if t == 'p':
            p = P('Box Text' if in_box else 'Normal')
            self.add_runs(p, b['r'], box_dark=dark)
            if b.get('lead') and p.runs:
                p.runs[0].font.size = Pt(10.5)
            return

    def build(self):
        d = self.doc
        cover = IR[0]
        img = next(b for b in cover['b'] if b['t'] == 'img')
        p = d.paragraphs[0] if d.paragraphs else d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(os.path.join(IMG, img['id'] + '.png'), width=Mm(210), height=Mm(296))
        # 본문 구역
        s = d.add_section(WD_SECTION.NEW_PAGE)
        s.left_margin = s.right_margin = Mm(20); s.top_margin = Mm(20); s.bottom_margin = Mm(18)
        s.header_distance = Mm(9); s.footer_distance = Mm(8)
        s.header.is_linked_to_previous = False; s.footer.is_linked_to_previous = False
        hp = s.header.paragraphs[0]; hp.style = d.styles['Small']
        r = hp.add_run('Heurēsis'); r.bold = True; set_font(r, DISPLAY, 9, color=INK)
        hp.add_run('  Vol.01 창간특집호 · 2026.10')
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; fp.style = d.styles['Small']
        run = fp.add_run()
        for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
            if kind:
                el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind); run._r.append(el)
            else:
                it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; run._r.append(it)

        # 차례  잡지 쪽 번호 대신 섹션 목록
        self.para(d, 'Kicker', text='IN THIS ISSUE')
        d.add_paragraph('차례', style='Heading 1')
        for pg in IR[3:]:
            if not pg['sec']: continue
            lab = next((''.join(r['t'] for r in b['r']) for b in pg['b'] if b['t'] == 'label'), '')
            head = next((''.join(r['t'] for r in b['r']).replace('\n', ' ') for b in pg['b'] if b['t'] == 'h'), '')
            lab = re.sub(r'\s*·\s*\d+\s*/\s*\d+\s*$', '', lab)
            p = d.add_paragraph(style='Normal'); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r1 = p.add_run(lab.upper() + '\n'); set_font(r1, MONO, 7.5, color=ACCENT[''])
            r2 = p.add_run(head); set_font(r2, SERIF, 12, True, INK)
        self.para(d, 'Small', text='이 문서는 인쇄용 잡지를 편집 가능한 흐름 문서로 옮긴 판이다  차트와 도식은 그림으로 넣었고 잡지의 쪽 번호는 쓰지 않는다  원래 지면은 PDF와 HTML 판에서 볼 수 있다')

        first = True
        for pg in IR[3:]:
            self.theme = pg['theme']
            if pg['sec'] or first:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first = False
            for b in pg['b']:
                self.block(d, b)
        os.makedirs(os.path.dirname(OUT), exist_ok=True)
        d.core_properties.title = 'Heurēsis Vol.01 창간특집호'
        d.core_properties.author = 'Heurēsis 편집부 · 편집장 McKay'
        d.core_properties.subject = '수학교육의 연구와 실천 사이'
        d.save(OUT)
        print(OUT, round(os.path.getsize(OUT) / 1e6, 1), 'MB')


if __name__ == '__main__':
    Builder().build()
