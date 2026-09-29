# -*- coding: utf-8 -*-
"""고천중학교 학생자치 개선안 DOCX 생성기

사용법 (프로젝트 폴더에서)
    python3 build/build_docx.py
    python3 build/build_docx.py --figs <그림 폴더> --out <출력 파일>

content.py 하나를 원본으로 삼아 out/고천중_학생자치_개선안.docx 를 만든다
"""
import argparse
import copy
import os
import re
import sys

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from content import META, COLORS, TOC, SUMMARY, BLOCKS  # noqa: E402

FONT = "맑은 고딕"

# 페이지 기하
PAGE_W, PAGE_H = 210, 297
MARGIN_LR, MARGIN_TB = 20, 18
CONTENT_W_MM = PAGE_W - 2 * MARGIN_LR          # 170mm
CONTENT_W_DXA = int(CONTENT_W_MM / 25.4 * 1440)  # twips
FIG_MAX_H_MM = 120                              # 그림 최대 높이 (한 쪽 안에 들어가도록)

TONES = {"ink": "ink", "del": "del", "raon": "raon", "lead": "lead", "ok": "ok"}
TINT = {"ink": "soft", "del": "del_l", "raon": "raon_l", "lead": "lead_l", "ok": "ok_l"}
ROW_ALT = "#F4F6FA"


def hx(key_or_hex):
    v = COLORS.get(key_or_hex, key_or_hex)
    return v.lstrip("#").upper()


def rgb(key_or_hex):
    return RGBColor.from_string(hx(key_or_hex))


# ───────────────────────────────── XML 도우미
def set_run_font(run, size=None, bold=None, color=None, italic=None):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)
    return run


def add_rich(par, text, size=None, color=None, bold_color=None, base_bold=False):
    """**굵게** 표기를 굵은 런으로, \\n 을 줄바꿈으로 바꾸어 넣는다"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**") and len(part) > 4
        seg = part[2:-2] if bold else part
        lines = seg.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                r = par.add_run()
                r.add_break()
                set_run_font(r, size=size)
            if line == "":
                continue
            r = par.add_run(line)
            set_run_font(r, size=size, bold=True if (bold or base_bold) else None,
                         color=(bold_color or color) if bold else color)
    return par


def p_border(par, side, color, sz=4, space=1, val="single"):
    ppr = par._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        # pBdr 는 spacing/ind 보다 앞에 와야 한다
        anchor = None
        for tag in ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc"):
            anchor = ppr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(bdr)
        else:
            ppr.append(bdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), hx(color))
    bdr.append(el)


def fmt(par, before=0, after=0, line=None, align=None, keep_next=None, left=None,
        first=None, keep_lines=None, rule=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        if rule == "exact":
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(line)
        else:
            pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if keep_lines is not None:
        pf.keep_together = keep_lines
    if left is not None:
        pf.left_indent = Pt(left)
    if first is not None:
        pf.first_line_indent = Pt(first)
    return par


def cell_shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hx(color))


def cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (sz, color) 또는 None(없음)"""
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcw = tcpr.find(qn("w:tcW"))
        if tcw is not None:
            tcw.addnext(b)
        else:
            tcpr.insert(0, b)
    for side in ("top", "left", "bottom", "right"):
        if side not in sides:
            continue
        el = OxmlElement(f"w:{side}")
        v = sides[side]
        if v is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(v[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), hx(v[1]))
        b.append(el)


def cell_margins(cell, top=None, bottom=None, left=None, right=None):
    tcpr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if v is None:
            continue
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tcpr.append(m)


def table_setup(table, widths_pct, border_color="line", cell_mar=(90, 150, 90, 150), borders=True):
    """고정 레이아웃, 폭, 테두리, 셀 여백 설정"""
    tbl = table._tbl
    tblpr = tbl.tblPr
    total = sum(widths_pct)
    dxa = [int(CONTENT_W_DXA * w / total) for w in widths_pct]
    dxa[-1] = CONTENT_W_DXA - sum(dxa[:-1])

    # 표 너비
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_W_DXA))
    tblw.set(qn("w:type"), "dxa")

    # 순서를 지키기 위해 기존 요소 제거 후 재배치
    for tag in ("w:tblBorders", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:jc", "w:tblInd"):
        for e in tblpr.findall(qn(tag)):
            tblpr.remove(e)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.append(jc)
    if borders:
        tb = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), hx(border_color))
            tb.append(el)
        tblpr.append(tb)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in zip(("top", "left", "bottom", "right"), cell_mar):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)
    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "0000")
    look.set(qn("w:firstRow"), "0")
    look.set(qn("w:lastRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "1")
    look.set(qn("w:noVBand"), "1")
    tblpr.append(look)

    # 그리드와 셀 폭
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(dxa[i]))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(dxa[i]))
            tcw.set(qn("w:type"), "dxa")
    return dxa


def row_props(row, header=False, cant_split=True):
    trpr = row._tr.get_or_add_trPr()
    if cant_split:
        trpr.append(OxmlElement("w:cantSplit"))
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def add_field(par, instr, size, color):
    r1 = par.add_run()
    set_run_font(r1, size=size, color=color)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._element.append(fc)
    r2 = par.add_run()
    set_run_font(r2, size=size, color=color)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._element.append(it)
    r3 = par.add_run()
    set_run_font(r3, size=size, color=color)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r3._element.append(fc)
    r4 = par.add_run("1")
    set_run_font(r4, size=size, color=color)
    r5 = par.add_run()
    set_run_font(r5, size=size, color=color)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r5._element.append(fc)


def inline_to_page_anchor(run):
    """인라인 그림을 쪽 왼쪽 위(0,0)에 고정한 글 뒤 그림으로 바꾼다 (표지 전용)"""
    inline = run._element.find(".//" + qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    docpr = inline.find(qn("wp:docPr"))
    frame = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = OxmlElement("wp:anchor")
    for k, v in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
                 ("simplePos", "0"), ("relativeHeight", "251658240"), ("behindDoc", "1"),
                 ("locked", "1"), ("layoutInCell", "1"), ("allowOverlap", "1")):
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    anchor.append(sp)
    for tag, rel in (("wp:positionH", "page"), ("wp:positionV", "page")):
        el = OxmlElement(tag)
        el.set("relativeFrom", rel)
        off = OxmlElement("wp:posOffset")
        off.text = "0"
        el.append(off)
        anchor.append(el)
    anchor.append(extent)
    ee = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        ee.set(k, "0")
    anchor.append(ee)
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(docpr)
    if frame is not None:
        anchor.append(frame)
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


# ───────────────────────────────── 스타일
def fix_style_fonts(style, size=None, bold=None, color=None):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in list(rfonts.attrib):
        if a.endswith("Theme"):
            del rfonts.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = rgb(color)
    style.font.italic = False


def setup_styles(doc):
    # 문서 기본값
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rf = rpr_default.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr_default.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    lang = rpr_default.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr_default.append(lang)
    lang.set(qn("w:val"), "ko-KR")
    lang.set(qn("w:eastAsia"), "ko-KR")

    for s in doc.styles:
        try:
            if s.type == 1 or s.type == 3:  # 문단, 표 스타일
                rpr = s.element.find(qn("w:rPr"))
                if rpr is not None and rpr.find(qn("w:rFonts")) is not None:
                    fix_style_fonts(s)
        except Exception:
            pass

    normal = doc.styles["Normal"]
    fix_style_fonts(normal, size=10.5, color="text")
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    # 한글 줄바꿈: 단어 단위
    ppr = normal.element.get_or_add_pPr()
    for tag in ("w:wordWrap", "w:autoSpaceDE", "w:autoSpaceDN"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "0")
        sp = ppr.find(qn("w:spacing"))
        if sp is not None:
            sp.addprevious(el)
        else:
            ppr.append(el)

    h1 = doc.styles["Heading 1"]
    fix_style_fonts(h1, size=20, bold=True, color="ink")
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.line_spacing = 1.1

    h2 = doc.styles["Heading 2"]
    fix_style_fonts(h2, size=13, bold=True, color="ink")
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.line_spacing = 1.15

    for name in ("Header", "Footer"):
        try:
            fix_style_fonts(doc.styles[name], size=8.5, color="muted")
        except KeyError:
            pass


# ───────────────────────────────── 블록 렌더러
class Builder:
    def __init__(self, figs_dir):
        self.doc = Document()
        self.figs = figs_dir
        self.missing = []
        setup_styles(self.doc)
        self.body = self.doc

    # 표지
    def cover(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(PAGE_W), Mm(PAGE_H)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(0)
        sec.header_distance = sec.footer_distance = Mm(0)
        sec.gutter = Mm(0)
        par = doc.add_paragraph()
        fmt(par, 0, 0, line=1, rule="exact")
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cover = os.path.join(self.figs, "cover.png")
        if os.path.exists(cover):
            r = par.add_run()
            set_run_font(r, size=1)
            r.add_picture(cover, width=Mm(PAGE_W), height=Mm(PAGE_H))
            inline_to_page_anchor(r)
        else:
            self.missing.append("cover.png")
            r = par.add_run(META["school"] + "  " + META["title"])
            set_run_font(r, size=24, bold=True, color="ink")
        # 문단 기호 크기를 최소화
        rpr = par._p.get_or_add_pPr()
        mrpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "2")
        mrpr.append(sz)
        rpr.append(mrpr)
        self.cover_par = par

        # 본문 구역
        new = doc.add_section(WD_SECTION.NEW_PAGE)
        # add_section 이 만든 빈 문단의 구역 정보를 표지 문단으로 옮겨 빈 쪽을 막는다
        body = doc.element.body
        holder = None
        for p in body.iterchildren(qn("w:p")):
            ppr = p.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                holder = p
        sectpr = holder.find(qn("w:pPr")).find(qn("w:sectPr"))
        if holder is not par._p:
            cpr = par._p.get_or_add_pPr()
            # sectPr 은 pPr 의 rPr 뒤에 위치
            cpr.append(sectpr)
            body.remove(holder)
        # 표지 구역은 머리말 꼬리말 없음 (본문 구역과 연결 해제)
        new.page_width, new.page_height = Mm(PAGE_W), Mm(PAGE_H)
        new.left_margin = new.right_margin = Mm(MARGIN_LR)
        new.top_margin = new.bottom_margin = Mm(MARGIN_TB)
        new.header_distance = Mm(9)
        new.footer_distance = Mm(8)
        # 쪽 번호 1부터
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        cols = new._sectPr.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(pg)
        else:
            new._sectPr.append(pg)
        self.header_footer(new)

    def header_footer(self, sec):
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        fmt(hp, 0, 0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        r = hp.add_run(f"{META['year']} 학생자치 개선안")
        set_run_font(r, size=8, color="muted")

        fp = sec.footer.paragraphs[0]
        fmt(fp, 0, 0)
        p_border(fp, "top", "line", sz=4, space=6)
        tabs = fp.paragraph_format.tab_stops
        # Footer 스타일의 기본 가운데/오른쪽 탭을 지우고 본문 폭 끝에 오른쪽 탭만 둔다
        for tab_twips in (4680, 9360):
            tabs.add_tab_stop(Emu(tab_twips * 635), WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Mm(CONTENT_W_MM), WD_TAB_ALIGNMENT.RIGHT)
        r = fp.add_run(f"{META['school']} {META['title']}")
        set_run_font(r, size=8.5, color="muted")
        r = fp.add_run("\t")
        set_run_font(r, size=8.5)
        add_field(fp, "PAGE", 9, "ink")
        # 첫 구역(표지)은 머리말 꼬리말 참조를 두지 않는다
        s0 = self.doc.sections[0]._sectPr
        for tag in ("w:headerReference", "w:footerReference"):
            for e in s0.findall(qn(tag)):
                s0.remove(e)

    # 목차와 핵심 요약
    def front(self):
        doc = self.doc
        t = doc.add_paragraph()
        fmt(t, 0, 2)
        r = t.add_run("CONTENTS")
        set_run_font(r, size=9, bold=True, color="del")
        t = doc.add_paragraph()
        fmt(t, 0, 10)
        p_border(t, "bottom", "ink", sz=12, space=6)
        r = t.add_run("목차")
        set_run_font(r, size=22, bold=True, color="ink")

        for num, title in TOC:
            p = doc.add_paragraph()
            fmt(p, 0, 0, line=1.0)
            p.paragraph_format.space_before = Pt(5.5)
            p.paragraph_format.space_after = Pt(5.5)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(16))
            p_border(p, "bottom", "line", sz=4, space=5)
            r = p.add_run(num)
            set_run_font(r, size=12, bold=True, color="del")
            r = p.add_run("\t" + title)
            set_run_font(r, size=12, color="ink")

        sp = doc.add_paragraph()
        fmt(sp, 0, 0)
        sp.paragraph_format.space_before = Pt(22)
        r = sp.add_run("SUMMARY")
        set_run_font(r, size=9, bold=True, color="del")
        t = doc.add_paragraph()
        fmt(t, 0, 10, keep_next=True)
        p_border(t, "bottom", "ink", sz=12, space=6)
        r = t.add_run("핵심 요약")
        set_run_font(r, size=22, bold=True, color="ink")

        table = doc.add_table(rows=len(SUMMARY), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_setup(table, [22, 78], cell_mar=(110, 170, 110, 170))
        for i, (label, text) in enumerate(SUMMARY):
            tone = "lead" if label == "회장단" else "del"
            row = table.rows[i]
            row_props(row)
            c0, c1 = row.cells
            cell_shade(c0, TINT[tone])
            c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c0.paragraphs[0]
            fmt(p, 0, 0, line=1.3)
            r = p.add_run(label)
            set_run_font(r, size=10.5, bold=True, color=tone)
            p = c1.paragraphs[0]
            fmt(p, 0, 0, line=1.45)
            add_rich(p, text, size=10, color="text", bold_color="ink")

    # 장 제목
    def h1(self, num, title, lead):
        doc = self.doc
        p = doc.add_paragraph(style="Heading 1")
        fmt(p, 0, 4, line=1.05, keep_next=True)
        p.paragraph_format.page_break_before = True
        r = p.add_run(num)
        set_run_font(r, size=28 if len(num) <= 2 else 22, bold=True, color="del")
        r.add_break()
        r = p.add_run(title)
        set_run_font(r, size=20, bold=True, color="ink")
        p2 = doc.add_paragraph()
        fmt(p2, 2, 18, line=1.4, keep_next=True)
        p_border(p2, "bottom", "del", sz=12, space=10)
        add_rich(p2, lead, size=11, color="muted")

    def h2(self, title):
        p = self.doc.add_paragraph(style="Heading 2")
        fmt(p, 16, 8, line=1.15, keep_next=True, left=7)
        p_border(p, "left", "del", sz=24, space=6)
        r = p.add_run(title)
        set_run_font(r, size=13, bold=True, color="ink")

    def para(self, text):
        p = self.doc.add_paragraph()
        fmt(p, 0, 6, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_rich(p, text, size=10.5, color="text", bold_color="ink")

    def bullets(self, items):
        n = len(items)
        for i, it in enumerate(items):
            p = self.doc.add_paragraph()
            fmt(p, 0, 4 if i < n - 1 else 8, line=1.45, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                left=16, first=-11)
            p.paragraph_format.tab_stops.add_tab_stop(Pt(16))
            r = p.add_run("•\t")
            set_run_font(r, size=10.5, bold=True, color="del")
            add_rich(p, it, size=10.5, color="text", bold_color="ink")

    def spacer(self, pt=8):
        p = self.doc.add_paragraph()
        fmt(p, 0, 0, line=pt, rule="exact")
        r = p.add_run("")
        set_run_font(r, size=2)
        rpr = p._p.get_or_add_pPr()
        mr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "4")
        mr.append(sz)
        rpr.append(mr)

    def table(self, spec):
        cols, rows, widths = spec["cols"], spec["rows"], spec["widths"]
        tone = spec.get("tone", "ink")
        ncol = len(cols)
        table = self.doc.add_table(rows=1 + len(rows), cols=ncol)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_setup(table, widths)

        if tone == "split":
            head_colors = ["ink", "del", "raon"] + ["ink"] * (ncol - 3)
        elif tone == "split2":
            head_colors = ["del", "raon"] + ["ink"] * (ncol - 2)
        else:
            head_colors = [TONES.get(tone, "ink")] * ncol

        # 짧은 값만 있는 좁은 열은 가운데 정렬
        centered = []
        for j in range(ncol):
            vals = [r[j] for r in rows]
            short = all(len(v.strip()) <= 6 for v in vals)
            centered.append(short and widths[j] <= 20 and j > 0 or (j == 0 and all(len(v) <= 3 for v in vals)))

        first_bold = tone != "split2"

        hdr = table.rows[0]
        row_props(hdr, header=True)
        for j, c in enumerate(hdr.cells):
            cell_shade(c, head_colors[j])
            cell_borders(c, top=(4, head_colors[j]), bottom=(4, head_colors[j]))
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            fmt(p, 0, 0, line=1.2, keep_next=True,
                align=WD_ALIGN_PARAGRAPH.CENTER if centered[j] else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(cols[j])
            set_run_font(r, size=9.5, bold=True, color="#FFFFFF")

        for i, rv in enumerate(rows):
            row = table.rows[i + 1]
            row_props(row)
            fill = ROW_ALT if i % 2 == 0 else "#FFFFFF"
            for j, c in enumerate(row.cells):
                cell_shade(c, fill)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = c.paragraphs[0]
                txt = rv[j]
                indent = len(txt) - len(txt.lstrip(" "))
                txt = txt.strip(" ")
                fmt(p, 0, 0, line=1.35,
                    align=WD_ALIGN_PARAGRAPH.CENTER if centered[j] else WD_ALIGN_PARAGRAPH.LEFT)
                if j == 0 and indent:
                    p.paragraph_format.left_indent = Pt(10)
                    r = p.add_run("└ ")
                    set_run_font(r, size=9, color="muted")
                    add_rich(p, txt, size=9.5, color="muted")
                elif j == 0 and first_bold:
                    add_rich(p, txt, size=9.5, color="ink", base_bold=True)
                else:
                    add_rich(p, txt, size=9.5, color="text", bold_color="ink")

        if spec.get("note"):
            p = self.doc.add_paragraph()
            fmt(p, 4, 10, line=1.35)
            add_rich(p, "※ " + spec["note"], size=8.5, color="muted")
        else:
            self.spacer(10)

    def callout(self, tone, title, text):
        tone = tone if tone in TINT else "ink"
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_setup(table, [100], borders=False, cell_mar=(150, 240, 150, 220))
        row = table.rows[0]
        row_props(row)
        c = row.cells[0]
        cell_shade(c, TINT[tone])
        cell_borders(c, left=(36, tone), top=None, bottom=None, right=None)
        p = c.paragraphs[0]
        fmt(p, 0, 3, line=1.2, keep_next=True)
        r = p.add_run(title)
        set_run_font(r, size=10.5, bold=True, color=tone)
        p = c.add_paragraph()
        fmt(p, 0, 0, line=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        add_rich(p, text, size=10.5, color="text", bold_color=tone if tone != "ink" else "ink")
        self.spacer(12)

    def figure(self, fid, caption):
        path = os.path.join(self.figs, fid + ".png")
        p = self.doc.add_paragraph()
        fmt(p, 6, 4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        if os.path.exists(path):
            with Image.open(path) as im:
                w_px, h_px = im.size
            w_mm = CONTENT_W_MM
            h_mm = w_mm * h_px / w_px
            if h_mm > FIG_MAX_H_MM:
                h_mm = FIG_MAX_H_MM
                w_mm = h_mm * w_px / h_px
            r = p.add_run()
            r.add_picture(path, width=Mm(w_mm), height=Mm(h_mm))
        else:
            self.missing.append(fid + ".png")
            r = p.add_run(f"[그림 파일 없음  {fid}]")
            set_run_font(r, size=9, color="muted")
        c = self.doc.add_paragraph()
        fmt(c, 0, 12, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = c.add_run(caption)
        set_run_font(r, size=9, color="muted")

    def render(self):
        self.cover()
        self.front()
        for b in BLOCKS:
            kind = b[0]
            if kind == "h1":
                self.h1(b[1], b[2], b[3])
            elif kind == "h2":
                self.h2(b[1])
            elif kind == "p":
                self.para(b[1])
            elif kind == "bullets":
                self.bullets(b[1])
            elif kind == "table":
                self.table(b[1])
            elif kind == "callout":
                self.callout(b[1], b[2], b[3])
            elif kind == "figure":
                self.figure(b[1], b[2])
            elif kind == "pagebreak":
                # 장 제목이 새 쪽에서 시작하므로 별도 쪽 나눔을 넣지 않는다 (빈 쪽 방지)
                pass
            else:
                raise ValueError(f"알 수 없는 블록 {kind}")
        self.cleanup()

    def cleanup(self):
        # 문서 끝의 빈 문단 여백 최소화 (마지막 쪽 뒤 빈 쪽 방지)
        body = self.doc.element.body
        last = body.findall(qn("w:p"))[-1]
        from docx.text.paragraph import Paragraph
        lp = Paragraph(last, self.doc)
        if not lp.text.strip() and not last.findall(".//" + qn("w:drawing")):
            fmt(lp, 0, 0, line=1, rule="exact")

        # 기본 서식 파일의 zoom 요소에 필수 속성 보완
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

        cp = self.doc.core_properties
        cp.title = f"{META['school']} {META['title']}"
        cp.subject = META["subtitle"]
        cp.author = f"{META['school']} {META['dept']}"
        cp.last_modified_by = f"{META['school']} {META['dept']}"
        cp.keywords = ", ".join(META["keywords"])
        cp.category = META["year"]
        cp.comments = f"{META['year']} {META['title']}  {META['date']}"
        cp.language = "ko-KR"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--figs", default=os.path.join(ROOT, "figs"))
    ap.add_argument("--out", default=os.path.join(ROOT, "out", "고천중_학생자치_개선안.docx"))
    args = ap.parse_args()
    b = Builder(args.figs)
    b.render()
    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    b.doc.save(args.out)
    print("저장", args.out)
    if b.missing:
        print("경고  그림 파일 없음:", ", ".join(b.missing))


if __name__ == "__main__":
    main()
