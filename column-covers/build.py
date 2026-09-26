"""AI Insight Column 표지 생성기

사용법
  python build.py next                              시리즈별 다음 번호 보기
  python build.py build --src 원고폴더 --out 결과폴더   원고폴더에 있는 칼럼의 표지를 만들고 원고 앞에 붙여 zip으로 묶기
  옵션 --only main-5 spinoff-3 ...   특정 표지만 / --covers-only 원고 없이 표지 PNG PDF만

필요 패키지 pymupdf python-docx playwright (브라우저는 설치된 Chromium 사용)
"""
import argparse, copy, glob, os, re, shutil, subprocess, sys, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import arts
from registry import COVERS
from series import SERIES

FONT_DIR = os.path.join(HERE, '.fonts')
FONT_CSS = ('https://fonts.googleapis.com/css2?family=Noto+Serif+KR:wght@400;700;900'
            '&family=Noto+Sans+KR:wght@300;400;500;700;900'
            '&family=Playfair+Display:ital,wght@0,700;0,900;1,400;1,700&family=JetBrains+Mono:wght@400;500')
A4 = (210, 297)


# ---------------------------------------------------------------- registry
def cover_id(cv):
    return f"{cv['series']}-{cv['n']}"


def resolve(cv):
    s = SERIES[cv['series']]
    return dict(cv, id=cover_id(cv), theme=s['theme'], cat=s['cat'], footer=s['footer'],
                num=s['num'](cv['n']), numlabel=s['numlabel'](cv['n']),
                kicker=cv.get('kicker', s['kicker']), art_fn=getattr(arts, cv['art']))


def check_registry():
    seen = set()
    for cv in COVERS:
        k = cover_id(cv)
        if k in seen:
            sys.exit(f'번호 중복 {k}')
        seen.add(k)
        if not hasattr(arts, cv['art']):
            sys.exit(f"arts.py에 {cv['art']} 함수가 없음")
        for field in ('title', 'sub', 'tags', 'meta'):
            text = ' '.join(cv[field]) if isinstance(cv[field], list) else cv[field]
            if '"' in text or '“' in text or '”' in text:
                sys.exit(f'{k} {field}에 쌍따옴표가 있음')


def next_numbers():
    out = {}
    for key in SERIES:
        ns = [cv['n'] for cv in COVERS if cv['series'] == key]
        out[key] = (max(ns) + 1) if ns else (0 if key == 'main' else 1)
    return out


# ---------------------------------------------------------------- environment
def ensure_fonts():
    css_path = os.path.join(FONT_DIR, 'fonts.css')
    if os.path.exists(css_path):
        return css_path
    os.makedirs(FONT_DIR, exist_ok=True)
    css = subprocess.run(['curl', '-sS', FONT_CSS], capture_output=True, text=True, check=True).stdout
    blocks = []
    for i, blk in enumerate(re.findall(r'@font-face\s*{[^}]*}', css)):
        url = re.search(r'url\((.*?)\)', blk).group(1)
        subprocess.run(['curl', '-sS', '-o', os.path.join(FONT_DIR, f'f{i}.ttf'), url], check=True)
        blocks.append(blk.replace(url, f'f{i}.ttf'))
    open(css_path, 'w').write('\n'.join(blocks))
    return css_path


def chromium_path():
    if os.environ.get('CHROMIUM'):
        return os.environ['CHROMIUM']
    hits = sorted(glob.glob('/opt/pw-browsers/chromium-*/chrome-linux/chrome'))
    return hits[-1] if hits else None


def page_size_mm(src):
    """원고 첫 페이지 크기를 mm로 돌려준다"""
    if src and src.lower().endswith('.pdf'):
        import pymupdf
        r = pymupdf.open(src)[0].rect
        return round(r.width * 25.4 / 72, 1), round(r.height * 25.4 / 72, 1)
    if src and src.lower().endswith('.docx'):
        import docx
        s = docx.Document(src).sections[0]
        return round(s.page_width / 36000, 1), round(s.page_height / 36000, 1)
    return A4


# ---------------------------------------------------------------- cover html
def html(cv, size, css_href):
    t = cv['theme']
    w, h = size
    title_len = max(len(x) for x in cv['title'])
    fs = 30 if title_len <= 9 else (26 if title_len <= 11 else 22)
    if h < 285:
        fs -= 1
    tags = ''.join(f'<span class="tag">{x}</span>' for x in cv['tags'])
    title = '<br>'.join(cv['title'])
    korean_kicker = any('가' <= ch <= '힣' for ch in cv['kicker'])
    return f'''<!doctype html><html><head><meta charset="utf-8">
<link rel="stylesheet" href="{css_href}">
<style>
@page {{ size: {w}mm {h}mm; margin: 0 }}
* {{ box-sizing: border-box; margin: 0; padding: 0 }}
html, body {{ width: {w}mm; height: {h}mm; background: {t['bg']}; color: {t['ink']}; }}
body {{ font-family: 'Noto Sans KR', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; overflow: hidden }}
.page {{ position: relative; width: {w}mm; height: {h}mm; padding: 16mm 17mm 14mm; display: flex; flex-direction: column }}
.grain {{ position: absolute; inset: 0; pointer-events: none; opacity: {'.06' if t['dark'] else '.05'};
  background-image: radial-gradient({t['ink']} .5px, transparent .6px); background-size: 3px 3px }}
.edge {{ position: absolute; left: 0; top: 0; bottom: 0; width: 5mm; background: {t['acc']} }}
.top {{ display: flex; justify-content: space-between; align-items: baseline; font-family: 'JetBrains Mono', monospace;
  font-size: 7.6pt; letter-spacing: .22em; text-transform: uppercase }}
.top b {{ font-weight: 500 }}
.rule {{ height: 1.4px; background: {t['ink']}; margin: 3mm 0 0 }}
.rule.thin {{ height: .6px; opacity: .5; margin-top: 1.2mm }}
.head {{ display: flex; align-items: flex-end; gap: 6mm; margin-top: 7mm }}
.num {{ font-family: 'Playfair Display', serif; font-weight: 900; font-size: 68pt; line-height: .78; color: {t['acc']};
  letter-spacing: -.02em }}
.numlabel {{ padding-bottom: 1mm }}
.numlabel .a {{ font-family: 'Playfair Display', 'Noto Serif KR', serif; font-style: italic; font-size: 13pt }}
.numlabel .b {{ font-size: 8.4pt; color: {t['mute']}; margin-top: 1mm; letter-spacing: .02em }}
.art {{ flex: 1; display: flex; align-items: center; justify-content: center; min-height: 0; padding: 4mm 0 }}
.art svg {{ width: 100%; height: 100%; max-height: 118mm; overflow: visible }}
.lbl {{ font-family: 'Noto Sans KR'; font-size: 13px; font-weight: 500; letter-spacing: .04em }}
.mono {{ font-family: 'JetBrains Mono'; font-size: 11px }}
.kicker {{ font-family: 'Playfair Display', 'Noto Serif KR', serif; font-style: italic; font-size: 12pt;
  color: {t['acc'] if not t['dark'] else t['acc2']}; margin-bottom: 3mm }}
.kicker.ko {{ font-style: normal; font-weight: 700; font-size: 10.5pt }}
.title {{ font-family: 'Noto Serif KR', serif; font-weight: 900; font-size: {fs}pt; line-height: 1.24; letter-spacing: -.035em;
  word-break: keep-all }}
.sub {{ margin-top: 5mm; font-size: 9.6pt; line-height: 1.65; max-width: 150mm; color: {t['ink']}; opacity: .82;
  word-break: keep-all; padding-left: 3.5mm; border-left: 2px solid {t['acc']} }}
.foot {{ margin-top: 8mm }}
.tags {{ display: flex; flex-wrap: wrap; gap: 1.6mm }}
.tag {{ font-size: 7.4pt; padding: .9mm 2.6mm; border: .8px solid {t['ink']}; border-radius: 10mm; opacity: .8 }}
.bottom {{ display: flex; justify-content: space-between; align-items: baseline; margin-top: 4mm; padding-top: 2.6mm;
  border-top: 1.4px solid {t['ink']}; font-size: 8pt }}
.bottom .s {{ font-family: 'Noto Serif KR'; font-weight: 700; font-size: 9.4pt }}
.bottom .m {{ font-family: 'JetBrains Mono'; font-size: 7pt; letter-spacing: .18em; color: {t['mute']} }}
</style></head><body><div class="page">
<div class="grain"></div><div class="edge"></div>
<div class="top"><b>AI Insight Column</b><span>{cv['cat']}</span></div>
<div class="rule"></div><div class="rule thin"></div>
<div class="head"><div class="num">{cv['num']}</div>
  <div class="numlabel"><div class="a">{cv['numlabel']}</div><div class="b">{' · '.join(cv['meta'])}</div></div></div>
<div class="art">{cv['art_fn'](t)}</div>
<div class="kicker{' ko' if korean_kicker else ''}">{cv['kicker']}</div>
<div class="title">{title}</div>
<div class="sub">{cv['sub']}</div>
<div class="foot"><div class="tags">{tags}</div>
<div class="bottom"><span class="s">{cv['footer']}</span><span class="m">AI · EDUCATION · 2026</span></div></div>
</div></body></html>'''


# ---------------------------------------------------------------- attach
def attach_pdf(src, cover_pdf, dst):
    import pymupdf
    d = pymupdf.open(src)
    cov = pymupdf.open(cover_pdf)
    r = d[0].rect
    new = pymupdf.open()
    pg = new.new_page(width=r.width, height=r.height)
    pg.show_pdf_page(pg.rect, cov, 0)
    new.insert_pdf(d)
    new.set_metadata(d.metadata)
    new.save(dst, garbage=3, deflate=True)


def attach_docx(src, cover_png, dst):
    """맨 앞에 여백 0짜리 구역을 만들고 표지 이미지를 한 쪽 가득 넣는다"""
    import docx
    from docx.shared import Emu
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document(src)
    body = d.element.body
    sec = d.sections[0]
    p = d.add_paragraph()
    p.add_run().add_picture(cover_png, width=Emu(sec.page_width), height=Emu(sec.page_height - 25400))
    body.remove(p._p)
    body.insert(0, p._p)
    pPr = p._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    for k, v in (('before', '0'), ('after', '0'), ('line', '240'), ('lineRule', 'auto')):
        sp.set(qn('w:' + k), v)
    pPr.append(sp)
    sect = copy.deepcopy(body.find(qn('w:sectPr')))
    for tag in ('w:headerReference', 'w:footerReference', 'w:titlePg', 'w:pgNumType'):
        for e in sect.findall(qn(tag)):
            sect.remove(e)
    m = sect.find(qn('w:pgMar'))
    for k in ('top', 'bottom', 'left', 'right', 'header', 'footer', 'gutter'):
        m.set(qn('w:' + k), '0')
    pPr.append(sect)
    d.save(dst)


# ---------------------------------------------------------------- main
def safe_name(s):
    return re.sub(r'[\\/:*?"<>|]', '', s).strip()


def build(args):
    check_registry()
    css = ensure_fonts()
    from playwright.sync_api import sync_playwright
    covers = [resolve(c) for c in COVERS]
    if args.only:
        covers = [c for c in covers if c['id'] in args.only]
    if not args.covers_only:
        covers = [c for c in covers if os.path.exists(os.path.join(args.src, c['src']))]
    if not covers:
        sys.exit('만들 표지가 없음 (registry의 src 이름과 원고 파일 이름을 확인)')

    work = os.path.join(args.out, '_work')
    col_dir = os.path.join(args.out, 'AI_Insight_Column_표지', '칼럼_표지포함')
    img_dir = os.path.join(args.out, 'AI_Insight_Column_표지', '표지_이미지')
    for d in (work, col_dir, img_dir):
        os.makedirs(d, exist_ok=True)
    shutil.copy(css, work)
    for f in glob.glob(os.path.join(FONT_DIR, '*.ttf')):
        if not os.path.exists(os.path.join(work, os.path.basename(f))):
            os.symlink(f, os.path.join(work, os.path.basename(f)))

    with sync_playwright() as p:
        exe = chromium_path()
        b = p.chromium.launch(**({'executable_path': exe} if exe else {}))
        for cv in covers:
            src = os.path.join(args.src, cv['src']) if args.src else None
            size = page_size_mm(src if src and os.path.exists(src) else None)
            fp = os.path.join(work, cv['id'] + '.html')
            open(fp, 'w').write(html(cv, size, 'fonts.css'))
            pg = b.new_page(viewport={'width': round(size[0] * 96 / 25.4), 'height': round(size[1] * 96 / 25.4)},
                            device_scale_factor=2)
            pg.goto('file://' + os.path.abspath(fp))
            pg.wait_for_timeout(400)
            png = os.path.join(work, cv['id'] + '.png')
            pdf = os.path.join(work, cv['id'] + '.pdf')
            pg.screenshot(path=png)
            pg.pdf(path=pdf, width=f'{size[0]}mm', height=f'{size[1]}mm', print_background=True)
            pg.close()
            label = safe_name(f"{cv['num']}_{' '.join(cv['title'])}")
            shutil.copy(png, os.path.join(img_dir, label + '.png'))
            if not args.covers_only:
                dst = os.path.join(col_dir, cv['src'].replace('..pdf', '.pdf'))
                (attach_pdf(src, pdf, dst) if src.lower().endswith('.pdf') else attach_docx(src, png, dst))
            print('ok', cv['id'], cv['num'], label)
        b.close()

    zpath = os.path.join(args.out, args.zip)
    root = os.path.join(args.out, 'AI_Insight_Column_표지')
    with zipfile.ZipFile(zpath, 'w', zipfile.ZIP_DEFLATED) as z:
        for r, _, fs in os.walk(root):
            for f in sorted(fs):
                full = os.path.join(r, f)
                z.write(full, os.path.relpath(full, args.out))
    print('zip', zpath)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest='cmd', required=True)
    sub.add_parser('next')
    bp = sub.add_parser('build')
    bp.add_argument('--src', default='')
    bp.add_argument('--out', required=True)
    bp.add_argument('--only', nargs='*')
    bp.add_argument('--covers-only', action='store_true')
    bp.add_argument('--zip', default='AI_Insight_Column_표지.zip')
    a = ap.parse_args()
    if a.cmd == 'next':
        check_registry()
        for k, n in next_numbers().items():
            print(k, SERIES[k]['num'](n))
    else:
        build(a)
