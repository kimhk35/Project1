#!/usr/bin/env python3
"""PRAXIS 창간특집호 DOCX 변환

_build/print.html 의 페이지 조각을 흐르는 문서로 옮긴다
도표는 render.js charts 가 만든 _build/charts/<figure id>.png 를 넣는다
"""
import os
import re

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

SRC = os.path.dirname(os.path.abspath(__file__))
B = os.path.join(SRC, '_build')
OUT = os.path.normpath(os.path.join(SRC, '..'))
NAME = 'PRAXIS_Vol01_창간특집호'

NAVY = RGBColor(0x14, 0x22, 0x3A)
ACC = RGBColor(0xC8, 0x41, 0x2A)
TEAL = RGBColor(0x2E, 0x6B, 0x66)
GREY = RGBColor(0x6B, 0x6F, 0x77)
SERIF = 'Batang'
SANS = 'Malgun Gothic'


def set_font(run, name=SERIF, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for k in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rfonts.set(qn(k), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)


def para_border(p, color='C8412A', side='top', sz=12):
    ppr = p._element.get_or_add_pPr()
    bdr = ppr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        ppr.append(bdr)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), '4')
    el.set(qn('w:color'), color)
    bdr.append(el)


def clean(t):
    return re.sub(r'\s+', ' ', t).strip()


class Writer:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(20)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
        st = self.doc.styles['Normal']
        st.font.name = SERIF
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.line_spacing = 1.35
        self.figs = 0
        self._footer()

    def _footer(self):
        sec = self.doc.sections[0]
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('PRAXIS 창간특집호 · 2026년 9월 · ')
        set_font(r, SANS, 8, color=GREY)
        r = p.add_run()
        f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
        f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
        r._element.append(f1); r._element.append(it); r._element.append(f2)
        set_font(r, SANS, 8, color=GREY)

    # ---------- inline ----------
    def runs(self, p, node, size=None, font=SERIF, color=None, bold=None):
        for ch in node.children if isinstance(node, Tag) else [node]:
            if isinstance(ch, NavigableString):
                t = re.sub(r'\s+', ' ', str(ch))
                if t:
                    set_font(p.add_run(t), font, size, bold, color)
            elif isinstance(ch, Tag):
                if ch.name in ('svg', 'style', 'script'):
                    continue
                if ch.name == 'br':
                    p.add_run().add_break()
                elif 'ev' in ch.get('class', []):
                    lv = int(ch.get('data-l', '0') or 0)
                    r = p.add_run(' ' + '■' * lv + '□' * (5 - lv) + ' ')
                    set_font(r, SANS, (size or 10) * .8, False, TEAL)
                    r = p.add_run(clean(ch.get_text()) + ' ')
                    set_font(r, SANS, (size or 10) * .8, False, TEAL)
                elif ch.name in ('b', 'strong'):
                    self.runs(p, ch, size, font, color, True)
                elif ch.name == 'em' or 'hl' in ch.get('class', []):
                    for r in self._collect(p, ch, size, font, color, bold):
                        r.font.highlight_color = 7  # yellow
                elif ch.name == 'sup':
                    r = p.add_run(ch.get_text()); set_font(r, font, size, bold, color); r.font.superscript = True
                elif ch.name == 'sub':
                    r = p.add_run(ch.get_text()); set_font(r, font, size, bold, color); r.font.subscript = True
                elif ch.name == 'a':
                    self.runs(p, ch, size, font, TEAL, bold)
                else:
                    self.runs(p, ch, size, font, color, bold)

    def _collect(self, p, node, size, font, color, bold):
        before = len(p.runs)
        self.runs(p, node, size, font, color, bold)
        return p.runs[before:]

    # ---------- blocks ----------
    def heading(self, text, level):
        sizes = {0: 26, 1: 20, 2: 13, 3: 11, 4: 10}
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10 if level > 1 else 4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(clean(text))
        set_font(r, SERIF if level <= 1 else SANS, sizes.get(level, 10), True, NAVY)
        # outline level for navigation pane and TOC
        ppr = p._element.get_or_add_pPr()
        ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), str(min(level, 8)))
        ppr.append(ol)
        return p

    def kicker(self, node):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(clean(node.get_text(' ')).upper())
        set_font(r, SANS, 8, True, ACC)

    def figure(self, node):
        fid = node.get('id')
        png = os.path.join(B, 'charts', f'{fid}.png') if fid else None
        if png and os.path.exists(png):
            from PIL import Image
            w, h = Image.open(png).size
            self.doc.add_picture(png, width=Mm(min(160, 105 * w / h)))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = node.find('figcaption')
        if cap:
            p = self.doc.add_paragraph()
            self.runs(p, cap, 8.5, SANS, GREY)

    def table(self, node):
        rows = node.find_all('tr')
        if not rows:
            return
        ncol = max(len(r.find_all(['td', 'th'])) for r in rows)
        t = self.doc.add_table(rows=0, cols=ncol)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r in rows:
            cells = r.find_all(['td', 'th'])
            row = t.add_row().cells
            for i, c in enumerate(cells):
                p = row[i].paragraphs[0]
                head = c.name == 'th'
                self.runs(p, c, 8.5, SANS, NAVY if head or 'k' in c.get('class', []) else None,
                          True if head or 'k' in c.get('class', []) else None)
                if head:
                    shade(row[i], 'EFE8DA')
        self.doc.add_paragraph()

    def boxed(self, node, fill='F3EEE4', title_color=ACC):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = 'Table Grid'
        tblpr = t._tbl.tblPr
        mar = OxmlElement('w:tblCellMar')
        for side in ('top', 'left', 'bottom', 'right'):
            e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), '140'); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblpr.append(mar)
        cell = t.rows[0].cells[0]
        shade(cell, fill)
        first = True
        dark = fill == '14223A'
        for ch in node.children:
            if not isinstance(ch, Tag):
                continue
            if ch.name == 'h4':
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                tag = ch.find(class_='tag')
                if tag:
                    r = p.add_run(clean(tag.get_text()) + '  '); set_font(r, SANS, 7.5, True, title_color)
                    tag.extract()
                r = p.add_run(clean(ch.get_text(' '))); set_font(r, SANS, 10, True, RGBColor(255, 255, 255) if dark else NAVY)
            else:
                self.block_into(cell, ch, dark)
            first = False
        self.doc.add_paragraph()

    def block_into(self, cell, node, dark=False):
        color = RGBColor(0xE6, 0xE1, 0xD6) if dark else None
        if node.name in ('ul', 'ol'):
            for i, li in enumerate(node.find_all('li', recursive=False), 1):
                p = cell.add_paragraph()
                mark = f'{i}. ' if node.name == 'ol' else '• '
                r = p.add_run(mark); set_font(r, SANS, 9, True, ACC)
                self.runs(p, li, 9, SERIF, color)
        elif node.name == 'table':
            for tr in node.find_all('tr'):
                p = cell.add_paragraph()
                cells = tr.find_all(['td', 'th'])
                for j, c in enumerate(cells):
                    r = p.add_run(clean(c.get_text(' ')) + ('  |  ' if j < len(cells) - 1 else ''))
                    set_font(r, SANS, 8.5, j == 0, color)
        elif node.name in ('p', 'div', 'dl', 'span', 'b', 'em', 'i'):
            if node.find(['p', 'ul', 'ol', 'table', 'div'], recursive=False) and node.name == 'div':
                for ch in node.children:
                    if isinstance(ch, Tag):
                        self.block_into(cell, ch, dark)
                return
            if node.name == 'dl':
                for dt in node.find_all('dt'):
                    dd = dt.find_next_sibling('dd')
                    p = cell.add_paragraph()
                    r = p.add_run(clean(dt.get_text()) + '  '); set_font(r, SANS, 8.5, True, ACC)
                    if dd:
                        self.runs(p, dd, 9, SERIF, color)
                return
            txt = clean(node.get_text(' '))
            if txt:
                p = cell.add_paragraph()
                self.runs(p, node, 9, SERIF, color)
        elif node.name == 'figure':
            pass

    def card(self, node):
        num = node.find(class_='num')
        h4 = node.find('h4')
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        para_border(p, '14223A', 'top', 8)
        if num:
            r = p.add_run(clean(num.get_text()) + '  '); set_font(r, SANS, 12, True, ACC)
        if h4:
            r = p.add_run(clean(h4.get_text(' '))); set_font(r, SERIF, 12, True, NAVY)
        meta = node.find(class_='meta')
        src = node.find(class_='src') or node.find(class_='who')
        for m in (meta, src):
            if m and clean(m.get_text()):
                q = self.doc.add_paragraph()
                self.runs(q, m, 8.5, SANS, GREY)
        for ch in node.children:
            if not isinstance(ch, Tag) or ch in (num, h4, meta, src):
                continue
            if ch.name == 'div' and 'top' in ch.get('class', []):
                sc = ch.find(class_='scores')
                if sc:
                    q = self.doc.add_paragraph()
                    self.runs(q, sc, 8.5, SANS, GREY)
                continue
            self.block(ch)
        self.doc.add_paragraph()

    def block(self, node):
        if not isinstance(node, Tag):
            t = clean(str(node))
            if t:
                p = self.doc.add_paragraph(); set_font(p.add_run(t), SERIF, 10)
            return
        cls = node.get('class', [])
        name = node.name
        if name in ('style', 'script', 'svg') or 'runhead' in cls or 'folio' in cls or 'cont' in cls:
            return
        if 'kicker' in cls:
            return self.kicker(node)
        if 'special-tag' in cls:
            p = self.doc.add_paragraph(); r = p.add_run(clean(node.get_text())); set_font(r, SANS, 9, True, ACC); return
        if name == 'h1':
            return self.heading(node.get_text(' '), 1)
        if name == 'h2':
            return self.heading(node.get_text(' '), 1)
        if name == 'h3':
            return self.heading(node.get_text(' '), 2)
        if name in ('h4', 'h5'):
            return self.heading(node.get_text(' '), 3)
        if 'cs-sub' in cls:
            p = self.doc.add_paragraph(); r = p.add_run(clean(node.get_text(' '))); set_font(r, SERIF, 14, True, ACC); return
        if 'dek' in cls:
            p = self.doc.add_paragraph(); self.runs(p, node, 11, SANS, RGBColor(0x3A, 0x3F, 0x48)); return
        if 'byline' in cls or 'note' in cls or 'cap' in cls or 'refs-mini' in cls:
            p = self.doc.add_paragraph(); self.runs(p, node, 8.5, SANS, GREY); return
        if 'pull' in cls:
            sm = node.find('small')
            if sm:
                sm.extract()
            p = self.doc.add_paragraph()
            para_border(p, 'C8412A', 'top', 16); para_border(p, 'CFC5B2', 'bottom', 6)
            self.runs(p, node, 13, SERIF, NAVY, True)
            if sm:
                q = self.doc.add_paragraph(); self.runs(q, sm, 8.5, SANS, GREY)
            return
        if name == 'figure':
            return self.figure(node)
        if name == 'table':
            return self.table(node)
        if 'box' in cls:
            fill = '14223A' if 'dark' in cls else 'F3EEE4'
            return self.boxed(node, fill)
        if 'card' in cls or 'app' in cls:
            return self.card(node)
        if name in ('ul', 'ol'):
            for i, li in enumerate(node.find_all('li', recursive=False), 1):
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(5)
                mark = f'{i}. ' if name == 'ol' else '• '
                r = p.add_run(mark); set_font(r, SANS, 10, True, ACC)
                self.runs(p, li, 10)
            return
        if name == 'dl':
            for dt in node.find_all('dt'):
                dd = dt.find_next_sibling('dd')
                p = self.doc.add_paragraph()
                r = p.add_run(clean(dt.get_text()) + '  '); set_font(r, SANS, 9, True, ACC)
                if dd:
                    self.runs(p, dd, 10)
            return
        if name == 'p':
            if not clean(node.get_text()):
                return
            p = self.doc.add_paragraph()
            size = 11 if 'lead' in cls else 10
            self.runs(p, node, size)
            return
        if name in ('div', 'aside', 'section', 'dd', 'figure', 'li', 'span', 'i', 'b', 'em', 'u', 'a'):
            has_block = any(isinstance(c, Tag) and c.name in ('p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'ul', 'ol', 'table',
                                                             'figure', 'dl', 'aside') for c in node.children)
            if has_block:
                for ch in node.children:
                    self.block(ch)
            else:
                t = clean(node.get_text(' '))
                if t:
                    p = self.doc.add_paragraph(); self.runs(p, node, 10)
            return

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc_field(self):
        p = self.doc.add_paragraph()
        r = p.add_run()
        f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'TOC \\o "1-2" \\h \\z \\u'
        f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
        t = OxmlElement('w:t'); t.text = '목차를 보려면 이 영역을 마우스 오른쪽 단추로 눌러 필드 업데이트를 선택하세요'
        f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
        for el in (f1, it, f2, t, f3):
            r._element.append(el)


def main():
    html = open(os.path.join(B, 'print.html'), encoding='utf-8').read()
    soup = BeautifulSoup(html, 'html.parser')
    w = Writer()
    pages = soup.find_all('section', class_='page')
    # cover image
    cover = os.path.join(B, 'charts', 'cover.png')
    if os.path.exists(cover):
        w.doc.add_picture(cover, width=Mm(170))
        w.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    w.page_break()
    # contents
    w.heading('차례', 1)
    p = w.doc.add_paragraph()
    set_font(p.add_run('이 DOCX판은 편집 가능한 흐름형 문서다 본문 속 면 번호는 PDF와 인쇄판의 쪽번호를 가리킨다'), SANS, 8.5, False, GREY)
    w.toc_field()
    for sec in pages:
        if sec.get('data-toc'):
            parts = sec['data-toc'].split('|') + ['', '']
            p = w.doc.add_paragraph()
            r = p.add_run(parts[0] + '  '); set_font(r, SANS, 8.5, True, ACC)
            r = p.add_run(parts[1]); set_font(r, SERIF, 10.5, True, NAVY)
            if parts[2]:
                r = p.add_run('  ' + parts[2]); set_font(r, SANS, 8.5, False, GREY)
    prev_sec = None
    for sec in pages[1:]:
        sid = sec.get('id', '')
        if sid in ('p-contents', 'p-contents2'):
            continue
        # new article → new page
        if sec.get('data-toc') or sec.get('data-bare') or prev_sec is None:
            w.page_break()
        prev_sec = sid
        for ch in sec.children:
            w.block(ch)
    out = os.path.join(OUT, f'{NAME}.docx')
    w.doc.save(out)
    print('docx written', out)


if __name__ == '__main__':
    main()
